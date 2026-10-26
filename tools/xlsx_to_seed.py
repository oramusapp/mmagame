#!/usr/bin/env python3
"""Konwertuje arkusz "thenotoriousg - SDCA" (.xlsx) na dane startowe aplikacji (JSON).

Użycie:  python3 -I tools/xlsx_to_seed.py <plik.xlsx> <wynik.json>
Potem:   SDCA_PASSWORD='...' node tools/encrypt_seed.mjs <wynik.json> seed.enc.js

Arkusz "SDCA v2" wyznacza aktywne wskaźniki. Pozostałe arkusze uzupełniają opisy,
historię Z-score, bazę wskaźników i bazę linków. "Submission Checklist" daje checklistę.
"""
import datetime
import json
import re
import sys
from urllib.parse import urlparse

import openpyxl

ACTIVE_SHEET = 'SDCA v2'
# kolejność: skąd brać nazwę/źródło oraz skąd opisy (v1.0 ma najpełniejsze opisy)
BASE_PRIORITY = ['SDCA v2', 'na tym pracuje', 'stare sdca v1.0', 'stare sdca v2.0']
DESC_PRIORITY = ['SDCA v2', 'stare sdca v1.0', 'na tym pracuje', 'stare sdca v2.0']
TEMPLATE_MARKERS = ('Write down pros, cons', 'normalized/logarithmic, Chart / Time Frame', 'if the RSI is below 20')
CAT_MAP = {'Fundamental Indicators': 'fundamental', 'Technical Indicators': 'technical', 'Sentiment Indicators': 'sentiment'}
CAT_NAMES = {'fundamental': 'Fundamental Indicators', 'technical': 'Technical Indicators', 'sentiment': 'Sentiment Indicators'}
GENERIC_LABELS = {'', 'Source', 'zasób', 'Alt / archiwum', 'resorces', 'websites for research', 'neotheescapist',
                  'Provided by', 'docs', 'google drive'}
SITE_NAMES = {
    'woocharts.com': 'Woocharts', 'lookintobitcoin.com': 'LookIntoBitcoin', 'checkonchain.com': 'Checkonchain',
    'woobull.com': 'Woobull Charts', 'web.archive.org': 'Web Archive (kopia)', 'buybitcoinworldwide.com': 'Buy Bitcoin Worldwide',
    'crypto-sentiment.com': 'Crypto Sentiment', 'tradingview.com': 'TradingView', 'chainexposed.com': 'ChainExposed',
    'augmento.ai': 'Augmento', 'colintalkscrypto.com': 'CBBI – ColinTalksCrypto', 'bybt.com': 'Bybt',
    'cane-island.digital': 'Cane Island Digital', 'quiverquant.com': 'Quiver Quant', 'coinglass.com': 'CoinGlass',
    'cryptoquant.com': 'CryptoQuant', 'bitbo.io': 'Bitbo Charts', 'chartinspect.com': 'ChartInspect',
    'sentimentrader.com': 'SentimenTrader', 'fidelity.com': 'Fidelity', 'alphractal.com': 'Alphractal',
    'kucoin.com': 'KuCoin', 'bluroo.ai': 'Bluroo', 'bitcoinmagazinepro.com': 'Bitcoin Magazine Pro',
    'capriole.com': 'Capriole', 'bgeometrics.com': 'BGeometrics', 'alternative.me': 'Alternative.me',
    'glassnode.com': 'Glassnode', 'velo.xyz': 'Velo', 'docs.google.com': 'Google Docs', 'drive.google.com': 'Google Drive',
    'pastebin.com': 'Pastebin', 'decentrader.com': 'Decentrader', 'mylongbow.com': 'Longbow',
}
DOC_HOSTS = ('docs.google.com', 'drive.google.com', 'pastebin.com')


def clean(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    return '' if any(s.startswith(m) for m in TEMPLATE_MARKERS) else s


def norm(name):
    s = re.sub(r'[\(\[].*?[\)\]]', '', name.lower())
    s = re.sub(r'[^a-z0-9]', '', s)
    return re.sub(r'^(the|bitcoin|btc)', '', s)


def url_key(u):
    return u.strip().rstrip('/').lower()


def is_url(s):
    return isinstance(s, str) and re.match(r'^(https?://|www\.)', s.strip(), re.I)


def cell_url(c):
    if c.hyperlink is not None and c.hyperlink.target:
        return c.hyperlink.target.strip()
    if is_url(c.value):
        return c.value.strip()
    return None


def parse_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, str):
        m = re.match(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', v.strip())
        if m:
            return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return None


def site_name(u):
    host = (urlparse(u).hostname or '').lower()
    host = host[4:] if host.startswith('www.') else host
    for k, v in SITE_NAMES.items():
        if host == k or host.endswith('.' + k):
            return v
    return host


def path_hint(u):
    p = urlparse(u)
    skip = {'charts', 'chart', 'pro', 'i', 'stats', 'script', 'edit', 'view', 'd', 'document', 'file', 'dashboard', 'market', 'news', 'insight'}
    segs = [s for s in p.path.split('/') if s and s.lower() not in skip]
    if 'chart=' in (p.query or ''):
        segs.append(re.search(r'chart=([^&]+)', p.query).group(1))
    for seg in reversed(segs):
        seg = re.sub(r'\.html?$', '', seg)
        if seg.isdigit() or len(seg) >= 30 or (len(seg) >= 16 and re.search(r'\d', seg) and '-' not in seg):
            continue
        seg = re.sub(r'^[A-Za-z0-9]{8}-', '', seg)  # id skryptu TradingView
        seg = re.sub(r'[-_]+', ' ', seg).strip()
        if seg:
            return seg[:1].upper() + seg[1:]
    return ''


def main(path, out):
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)

    rows, cat_desc, sheet_dates, raw_links = [], {}, {}, []
    for sh in BASE_PRIORITY:
        if sh not in wbf.sheetnames:
            continue
        wf, wv = wbf[sh], wbv[sh]
        sheet_dates[sh] = parse_date(wf['C1'].value)
        headers = {c: clean(wf.cell(3, c).value) for c in range(1, wf.max_column + 1)}
        cat = None
        for r in range(5, min(wf.max_row, 30) + 1):
            a = wf.cell(r, 1).value
            if isinstance(a, str) and a.strip() in CAT_MAP:
                cat = CAT_MAP[a.strip()]
            elif isinstance(a, str) and cat and len(a.strip()) > 40:
                cat_desc.setdefault(cat, a.strip())
            name = clean(wf.cell(r, 2).value)
            if not name or not cat:
                continue
            zv = wv.cell(r, 9).value
            src_cell = wf.cell(r, 10)
            source = cell_url(src_cell) or ''
            alt = cell_url(wf.cell(r, 11))
            alts = [alt] if alt and 'web.archive.org' in alt else []
            rows.append({
                'sheet': sh, 'row': r, 'category': cat, 'name': name,
                'why': clean(wf.cell(r, 4).value), 'how': clean(wf.cell(r, 5).value), 'score': clean(wf.cell(r, 6).value),
                'providedBy': clean(wf.cell(r, 7).value) or 'Select',
                'z': float(zv) if isinstance(zv, (int, float)) else None,
                'source': source, 'sourceText': '' if source else clean(src_cell.value), 'altLinks': alts,
            })
        for row in wf.iter_rows():
            for c in row:
                u = cell_url(c)
                if not u:
                    continue
                label = '' if is_url(c.value) else clean(c.value)
                is_ind_cell = c.column in (10, 11) and 5 <= c.row <= 30 and clean(wf.cell(c.row, 2).value)
                if is_ind_cell and not (c.column == 11 and 'web.archive.org' not in u):
                    continue
                raw_links.append({'url': u, 'sheet': sh, 'cell': c.coordinate, 'label': label, 'header': headers.get(c.column, '')})

    # ---- grupowanie tych samych wskaźników (ta sama nazwa albo ten sam link źródłowy)
    parent = list(range(len(rows)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    seen = {}
    for i, r in enumerate(rows):
        for key in ('n:' + norm(r['name']), 'u:' + url_key(r['source']) if r['source'] else None):
            if not key:
                continue
            if key in seen:
                parent[find(i)] = find(seen[key])
            else:
                seen[key] = i
    groups = {}
    for i in range(len(rows)):
        groups.setdefault(find(i), []).append(rows[i])

    def by(order, members):
        return sorted(members, key=lambda m: (order.index(m['sheet']), m['row']))

    indicators, sheet_ids = [], {}
    for members in groups.values():
        base = by(BASE_PRIORITY, members)[0]
        desc = by(DESC_PRIORITY, members)
        active = base['sheet'] == ACTIVE_SHEET
        gid = 'i' + str(len(indicators) + 1).zfill(3)
        pick = lambda f: next((m[f] for m in desc if m[f]), '')
        provided = next((m['providedBy'] for m in desc if m['providedBy'] not in ('', 'Select')), 'Select')
        if active:
            provided_note = provided if base['providedBy'] == 'Select' and provided != 'Select' else ''
        alt = []
        for m in by(BASE_PRIORITY, members):
            for u in ([m['source']] if m['source'] and m['source'] != base['source'] else []) + m['altLinks']:
                if url_key(u) not in {url_key(x) for x in alt}:
                    alt.append(u)
        history = sorted(
            [{'sheet': m['sheet'], 'date': sheet_dates.get(m['sheet']), 'z': m['z']} for m in members
             if m['z'] is not None and sheet_dates.get(m['sheet']) and m['sheet'] != ACTIVE_SHEET],
            key=lambda h: h['date'])
        z = base['z'] if active else (history[-1]['z'] if history else None)
        ind = {
            'id': gid, 'name': base['name'], 'category': base['category'], 'active': active,
            'order': base['row'] if active else 1000 + len(indicators),
            'z': z, 'why': pick('why'), 'how': pick('how'), 'score': pick('score'),
            'providedBy': base['providedBy'] if (active and base['providedBy'] != 'Select') else provided,
            'source': base['source'], 'sourceText': base['sourceText'], 'altLinks': alt,
            'tv': any('tradingview.com' in (m['source'] or '') for m in members) or 'tradingview' in base['sourceText'].lower(),
            'aliases': [{'sheet': m['sheet'], 'name': m['name']} for m in by(BASE_PRIORITY, members)],
            'descFrom': next((m['sheet'] for m in desc if m['why'] or m['how'] or m['score']), None),
            'history': history, 'updatedAt': None,
        }
        indicators.append(ind)
        for m in members:
            sheet_ids[(m['sheet'], m['row'])] = gid

    indicators.sort(key=lambda i: (not i['active'], ['fundamental', 'technical', 'sentiment'].index(i['category']), i['order']))
    for n, i in enumerate(indicators):
        i['order'] = n

    # ---- baza linków
    links = {}

    def add_link(u, origin, label='', header='', ind_id=None, cat=None):
        u = u.strip()
        if u.lower().startswith('www.'):
            u = 'https://' + u
        L = links.setdefault(url_key(u), {'url': u, 'origins': [], 'labels': [], 'indicators': [], 'cats': [], 'tags': []})
        if origin not in L['origins']:
            L['origins'].append(origin)
        if label and label not in GENERIC_LABELS and label not in L['labels']:
            L['labels'].append(label)
        if header == 'neotheescapist' and 'neotheescapist' not in L['tags']:
            L['tags'].append('neotheescapist')
        if ind_id and ind_id not in L['indicators']:
            L['indicators'].append(ind_id)
        if cat and cat not in L['cats']:
            L['cats'].append(cat)

    for i in indicators:
        for u in ([i['source']] if i['source'] else []) + i['altLinks']:
            add_link(u, 'wskaźnik', ind_id=i['id'], cat=i['category'])
    for l in raw_links:
        add_link(l['url'], f"{l['sheet']}!{l['cell']}", l['label'], l['header'])

    ind_by_id = {i['id']: i for i in indicators}
    link_list = []
    for n, L in enumerate(sorted(links.values(), key=lambda x: (not x['indicators'], site_name(x['url']).lower(), x['url']))):
        u = L['url']
        host = (urlparse(u).hostname or '').lower()
        site = site_name(u)
        tags = list(L['tags'])
        if 'web.archive.org' in host:
            tags.append('archiwum')
        if 'tradingview.com' in host:
            tags.append('TradingView')
        if L['indicators']:
            names = [ind_by_id[x]['name'] for x in L['indicators']]
            title = f"{names[0]} — {site}"
            cat = L['cats'][0]
        else:
            labels = [x for x in L['labels'] if x.lower() not in host]
            hint = labels[0] if labels else path_hint(u)
            if not hint and any(host.endswith(h) for h in DOC_HOSTS):
                hint = 'dokument'
            title = f"{site} — {hint}" if hint else site
            cat = 'docs' if any(host.endswith(h) for h in DOC_HOSTS) else 'research'
        link_list.append({'id': 'l' + str(n + 1).zfill(3), 'url': u, 'title': title, 'site': site, 'category': cat,
                          'tags': tags, 'indicators': L['indicators'], 'note': '',
                          'origins': [o for o in L['origins'] if o != 'wskaźnik']})

    # ---- historia aktualizacji z dawnych arkuszy (średnia jak w Excelu: AVERAGE kolumny Z)
    updates = []
    for sh in ('stare sdca v2.0', 'stare sdca v1.0'):
        vals = {sheet_ids[(r['sheet'], r['row'])]: r['z'] for r in rows if r['sheet'] == sh and r['z'] is not None}
        if not vals or not sheet_dates.get(sh):
            continue
        avg = sum(vals.values()) / len(vals)
        updates.append({'id': 'u-' + sh.replace(' ', '-').replace('.', ''), 'date': sheet_dates[sh], 'avgZ': round(avg, 4),
                        'count': len(vals), 'values': vals, 'note': f"Zaimportowane z arkusza „{sh}” ({len(vals)} wskaźników)"})
    updates.sort(key=lambda u: u['date'])

    # ---- checklista
    cl = wbf['Submission Checklist']
    checklist, section = [], None
    options = {}
    for dv in cl.data_validations.dataValidation:
        opts = [o.strip() for o in (dv.formula1 or '').strip('"').split(',')]
        for rng in str(dv.sqref).split():
            m = re.match(r'([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?', rng)
            a, b = int(m.group(2)), int(m.group(4) or m.group(2))
            for r in range(a, b + 1):
                options[r] = opts
    for r in range(6, cl.max_row + 1):
        b, d = cl.cell(r, 2).value, cl.cell(r, 4).value
        if not b:
            continue
        if d == 'STATUS':
            section = b.strip()
            continue
        text = b.strip()
        m = re.match(r'^(\d+)\.\s*(.*)$', text, re.S)
        checklist.append({'id': f'c{r}', 'section': section, 'num': int(m.group(1)) if m else None,
                          'text': m.group(2).strip() if m else text, 'status': d, 'options': options.get(r, [])})

    ft = wbf['Forward Testing'] if 'Forward Testing' in wbf.sheetnames else None
    seed = {
        'version': 1,
        'source': 'thenotoriousg - SDCA - aktualny do pracy.xlsx',
        'dateUpdated': sheet_dates.get(ACTIVE_SHEET),
        'categories': {k: {'name': CAT_NAMES[k], 'desc': cat_desc.get(k, '')} for k in CAT_NAMES},
        'indicators': indicators,
        'links': link_list,
        'updates': updates,
        'checklist': checklist,
        'notes': {'forward': clean(ft['C1'].value) if ft else ''},
    }
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(seed, f, ensure_ascii=False, indent=1)
    act = [i for i in indicators if i['active']]
    print(f"wskaźniki: {len(indicators)} (aktywne {len(act)}), linki: {len(link_list)}, aktualizacje: {len(updates)}, checklista: {len(checklist)}")


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
